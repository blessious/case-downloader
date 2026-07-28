from __future__ import annotations

import csv
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import FALLBACK_REPORT_DIR


CASE_HEADERS = [
    "Ordinal",
    "Week",
    "Assigned Title",
    "Filename",
    "G.R. Number",
    "Decision Date",
    "Final Status",
    "Saved Path",
    "Source",
    "Source URL",
    "Error Code",
    "Error Detail",
    "Warnings",
    "Updated At",
]
ATTEMPT_HEADERS = [
    "Case Ordinal",
    "Assigned Title",
    "Source",
    "URL",
    "Attempt",
    "Status",
    "Error Code",
    "Error Detail",
    "Started At",
    "Finished At",
]


def _case_row(item: dict[str, Any]) -> list[Any]:
    return [
        item["ordinal"],
        item["week"],
        item["title"],
        item["filename"],
        item["docket"],
        item["decision_date"],
        item["status"],
        item["saved_path"],
        item["source"],
        item["source_url"],
        item["error_code"],
        item["error_message"],
        "; ".join(item.get("warnings", [])),
        item["updated_at"],
    ]


def _style_sheet(sheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="153A5B")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        width = max(
            12,
            min(
                55,
                max(
                    len(str(sheet.cell(row=row, column=column).value or ""))
                    for row in range(1, min(sheet.max_row, 150) + 1)
                )
                + 2,
            ),
        )
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def create_xlsx_report(run: dict[str, Any], destination: Path) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Status", "Count"])
    counts = Counter(item["status"] for item in run["cases"])
    summary.append(["Total", len(run["cases"])])
    for status in ("Saved", "Already saved", "Not saved", "Duplicate input", "Cancelled"):
        summary.append([status, counts.get(status, 0)])
    summary.append([])
    summary.append(["Run ID", run["id"]])
    summary.append(["Created", run["created_at"]])
    summary.append(["Output Folder", run["output_dir"]])
    _style_sheet(summary)

    all_cases = workbook.create_sheet("All Cases")
    all_cases.append(CASE_HEADERS)
    for item in run["cases"]:
        all_cases.append(_case_row(item))
    _style_sheet(all_cases)

    not_saved = workbook.create_sheet("Not Saved")
    not_saved.append(CASE_HEADERS)
    for item in run["cases"]:
        if item["status"] not in ("Saved", "Already saved"):
            not_saved.append(_case_row(item))
    _style_sheet(not_saved)

    attempts = workbook.create_sheet("Attempts")
    attempts.append(ATTEMPT_HEADERS)
    cases_by_id = {item["id"]: item for item in run["cases"]}
    for attempt in run["attempts"]:
        case = cases_by_id.get(attempt["case_id"], {})
        attempts.append(
            [
                case.get("ordinal", ""),
                case.get("title", ""),
                attempt["source"],
                attempt["url"],
                attempt["attempt_no"],
                attempt["status"],
                attempt["error_code"],
                attempt["error_message"],
                attempt["started_at"],
                attempt["finished_at"],
            ]
        )
    _style_sheet(attempts)

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination


def create_fallback_reports(run: dict[str, Any], destination_dir: Path | None = None) -> list[Path]:
    target = destination_dir or FALLBACK_REPORT_DIR
    target.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = target / f"case-download-report_{timestamp}"
    json_path = base.with_suffix(".json")
    csv_path = base.with_suffix(".csv")
    json_path.write_text(json.dumps(run, indent=2, ensure_ascii=False), encoding="utf-8")
    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(CASE_HEADERS)
        for item in run["cases"]:
            writer.writerow(_case_row(item))
    return [csv_path, json_path]
