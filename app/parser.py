from __future__ import annotations

import io
import re
from datetime import datetime

from docx import Document
from pypdf import PdfReader

from .domain import ParsedCase
from .filenames import normalize_case_title


MONTHS = (
    "Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|"
    "Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|"
    "Nov(?:ember)?|Dec(?:ember)?"
)
DATE_RE = re.compile(
    rf"(?P<date>(?:(?:{MONTHS})\.?\s+\d{{1,2}},?\s+\d{{4}}|"
    rf"\d{{1,2}}\s+(?:{MONTHS})\.?\s+\d{{4}}))",
    re.IGNORECASE,
)
WEEK_RE = re.compile(r"^\s*WEEK\s*(?P<number>\d+)\s*:?\s*$", re.IGNORECASE)
SECTION_RE = re.compile(
    r"^\s*(?P<label>(?:MODULE|UNIT|TOPIC)\s+[A-Z0-9IVX]+|"
    r"(?:ASSIGNED\s+)?CASES?|(?:ASSIGNED\s+)?READINGS?)\s*:?\s*$",
    re.IGNORECASE,
)
LIST_PREFIX_RE = re.compile(r"^\s*(?:\d+[\.\)]|[-\u2022])\s+")
CASE_TITLE_RE = re.compile(r"(?<!\w)(?:v(?:s)?\.?|versus)(?!\w)", re.IGNORECASE)
DOCKET_START_RE = re.compile(
    r"(?<![A-Za-z])(?:G\.?\s*R\.?(?:\s*R\.?)?\s*(?:(?:Nos?\.?)|#)?|L-?\s*(?=\d{3,}))\s*",
    re.IGNORECASE,
)
DOCKET_VALUE_RE = re.compile(
    r"(?P<docket>(?:G\.?\s*R\.?(?:\s*R\.?)?\s*(?:(?:Nos?\.?)|#)?\s*[-:]?\s*)?"
    r"(?:L\s*-?\s*)?\d{3,}(?:\s*(?:,|&|and|/)\s*(?:L\s*-?\s*)?\d{3,})*)",
    re.IGNORECASE,
)
REPORTER_RE = re.compile(
    r"(?P<citation>\b\d{1,4}\s*(?:SCRA|PHIL\.?)\s*\d{1,4}\b)",
    re.IGNORECASE,
)
YEAR_ONLY_RE = re.compile(r"\(\s*(?:18|19|20)\d{2}\s*\)\s*[;,.]?$", re.IGNORECASE)


def parse_date(value: str) -> str:
    clean = re.sub(r"\s+", " ", value.strip().replace(",", "").replace(".", ""))
    for pattern in ("%B %d %Y", "%b %d %Y", "%d %B %Y", "%d %b %Y"):
        try:
            return datetime.strptime(clean, pattern).date().isoformat()
        except ValueError:
            continue
    return ""


def _is_heading(line: str) -> bool:
    return bool(WEEK_RE.match(line) or SECTION_RE.match(line))


def normalize_syllabus_line(value: str) -> str:
    """Repair common PDF text-extraction spacing without assuming one citation format."""
    line = re.sub(r"\s+", " ", value.strip())
    heavily_spaced = len(re.findall(r"\b[A-Za-z0-9]\b", line)) >= 8

    if heavily_spaced:
        previous = ""
        while previous != line:
            previous = line
            line = re.sub(
                r"(?<=[A-Za-z0-9])\s+(?=[A-Za-z0-9](?:\s|[.,);:-]))",
                "",
                line,
            )
        line = re.sub(r"\s*([.(),;:-])\s*", r"\1", line)
        line = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", line)
        line = re.sub(r"\b(G\.?\s*R\.?)\s*No\.?", r"\1 No. ", line, flags=re.IGNORECASE)

    def collapse_spaced_run(match: re.Match[str]) -> str:
        return re.sub(r"\s+", "", match.group(0))

    line = re.sub(
        r"(?<![A-Za-z0-9])(?:[A-Za-z0-9]\s+){3,}[A-Za-z0-9](?![A-Za-z0-9])",
        collapse_spaced_run,
        line,
    )
    month_pattern = r"Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?"
    line = re.sub(rf"(?<=\d)(?={month_pattern}\b)", " ", line, flags=re.IGNORECASE)
    line = re.sub(rf"\b({month_pattern})\.?\s*(?=\d)", r"\1 ", line, flags=re.IGNORECASE)
    line = re.sub(r",(?=\d{4}\b)", ", ", line)
    line = re.sub(r"\.(?=(?:v(?:s)?|versus)\.)", " ", line, flags=re.IGNORECASE)
    line = re.sub(r"(?<=\.)v(?:s)?\.(?=[A-Z])", " v. ", line, flags=re.IGNORECASE)
    line = re.sub(r"\bv(?:s)?\.(?=[A-Za-z])", "v. ", line, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", line).strip()


def _looks_like_case(line: str) -> bool:
    return bool(
        CASE_TITLE_RE.search(line)
        or DOCKET_START_RE.search(line)
        or REPORTER_RE.search(line)
        or re.match(r"^Case\b", line, re.IGNORECASE)
    )


def _logical_lines(text: str) -> list[str]:
    logical: list[str] = []
    current = ""
    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = normalize_syllabus_line(raw)
        if not line:
            if current:
                logical.append(current)
                current = ""
            continue
        if _is_heading(line):
            if current:
                logical.append(current)
                current = ""
            logical.append(line)
            continue

        has_prefix = bool(LIST_PREFIX_RE.match(line))
        stripped = LIST_PREFIX_RE.sub("", line, count=1).strip()
        starts_case = _looks_like_case(stripped)
        if has_prefix or (starts_case and CASE_TITLE_RE.search(stripped)):
            if current:
                logical.append(current)
            current = stripped
        elif current:
            current = f"{current} {stripped}"
        elif starts_case:
            current = stripped
    if current:
        logical.append(current)
    return logical


def _clean_metadata_edges(value: str) -> str:
    return value.strip().strip(" ,;()[]")


def _parse_case_line(line: str, ordinal: int, week: str) -> ParsedCase:
    warnings: list[str] = []
    decision_date = ""
    date_match = DATE_RE.search(line)
    if date_match:
        decision_date = parse_date(date_match.group("date"))
        if not decision_date:
            warnings.append(f"Unrecognized date: {date_match.group('date')}")

    reporter_match = REPORTER_RE.search(line)
    reporter_citation = reporter_match.group("citation") if reporter_match else ""
    reporter_citation = re.sub(
        r"(?i)(\d)(SCRA|PHIL\.?)", r"\1 \2", reporter_citation
    )
    reporter_citation = re.sub(
        r"(?i)(SCRA|PHIL\.?)(\d)", r"\1 \2", reporter_citation
    )
    reporter_citation = re.sub(r"\s+", " ", reporter_citation).strip()

    docket = ""
    docket_match = DOCKET_START_RE.search(line)
    if docket_match:
        value_match = DOCKET_VALUE_RE.match(line, docket_match.start())
        if value_match:
            docket = _clean_metadata_edges(value_match.group("docket"))

    year_match = YEAR_ONLY_RE.search(line)

    metadata_starts = [
        match.start()
        for match in (docket_match, reporter_match, date_match, year_match)
        if match is not None
    ]
    title = line[: min(metadata_starts)] if metadata_starts else line
    title = normalize_case_title(_clean_metadata_edges(title))
    title = re.sub(r"^(?:Cases?)\s*:\s*", "", title, flags=re.IGNORECASE)

    errors: list[str] = []
    if not title:
        errors.append("Missing case title")
    if not docket:
        warnings.append("Missing or malformed docket number")
    if not decision_date:
        warnings.append("Missing or malformed decision date")
    return ParsedCase(
        ordinal=ordinal,
        week=week,
        title=title,
        docket=docket,
        decision_date=decision_date,
        raw_text=line,
        parse_error="; ".join(errors),
        warnings=warnings,
        reporter_citation=reporter_citation,
    )


def parse_text(text: str) -> list[ParsedCase]:
    cases: list[ParsedCase] = []
    section = "Uncategorized"
    for line in _logical_lines(text):
        week_match = WEEK_RE.match(line)
        if week_match:
            section = f"Week {int(week_match.group('number'))}"
            continue
        section_match = SECTION_RE.match(line)
        if section_match:
            label = section_match.group("label")
            if re.match(r"(?:MODULE|UNIT|TOPIC)", label, re.IGNORECASE):
                section = re.sub(r"\s+", " ", label).title()
            continue
        if not _looks_like_case(line):
            continue
        cases.append(_parse_case_line(line, len(cases) + 1, section))
    return cases


def extract_docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" ".join(cells))
    return "\n".join(parts)


def parse_page_range(value: str, page_count: int) -> list[int]:
    if not value.strip():
        return list(range(page_count))
    selected: list[int] = []
    for part in value.split(","):
        token = part.strip()
        if not token:
            continue
        if "-" in token:
            start_text, end_text = token.split("-", 1)
            if not start_text.isdigit() or not end_text.isdigit():
                raise ValueError("Pages must look like 3, 3-6, or 3,5-7")
            start, end = int(start_text), int(end_text)
        elif token.isdigit():
            start = end = int(token)
        else:
            raise ValueError("Pages must look like 3, 3-6, or 3,5-7")
        if start < 1 or end < start or end > page_count:
            raise ValueError(f"Choose pages between 1 and {page_count}")
        selected.extend(range(start - 1, end))
    if not selected:
        raise ValueError("Enter at least one PDF page")
    return list(dict.fromkeys(selected))


def extract_pdf_text(data: bytes, page_range: str = "") -> str:
    reader = PdfReader(io.BytesIO(data), strict=False)
    parts: list[str] = []
    for page_index in parse_page_range(page_range, len(reader.pages)):
        page = reader.pages[page_index]
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n\n".join(parts)


def docket_tokens(docket: str) -> list[str]:
    tokens = re.findall(r"(?:L-)?\d{3,}", docket.upper())
    return [token.replace("L-", "") for token in tokens]
