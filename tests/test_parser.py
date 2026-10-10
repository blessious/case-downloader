from __future__ import annotations

import io

import pytest
from docx import Document

from app.parser import extract_docx_text, parse_page_range, parse_text


def test_gray_surname_is_not_a_docket_prefix():
    case = parse_text("M.E. Gray vs. Insular Lumber (67 Phil 139)")[0]
    assert case.title == "M.E. Gray v. Insular Lumber"
    assert case.docket == ""


def test_consolidated_space_separated_and_shorthand_dockets_preserved():
    case = parse_text("Lu Ym v. Lu Ym, G.R. Nos. 219902 219903 & 219943-44, 17 January 2018")[0]
    assert case.docket == "G.R. Nos. 219902 219903 & 219943-44"


SAMPLE = """ASSIGNED CASES

WEEK 2
1. Honda Philippines, Inc. v. Samahan ng Malalaya sa Honda, G.R. No. 145561, June 15, 2005
2. TMPCA v. NLRC, G.R. Nos. 158786 & 158789, October 19, 2007

WEEK 3
1. Duncans Association of Detailman - PTGWO vs Glaxo Wellcome, G.R. No. 162994, September 17, 2004
"""


@pytest.mark.parametrize(
    ("entry", "title", "docket", "date", "reporter"),
    [
        (
            "M i n d a n a o B u s . C o . v s . C i t y A s s e s s o r "
            "( G . R N o . L - 1 7 8 7 0 S e p t e m b e r 2 9 , 1 9 6 2 ) ;",
            "Mindanao Bus.Co v. City Assessor",
            "G.R No. L-17870",
            "1962-09-29",
            "",
        ),
        (
            "Mun. of Cavite v.R o j a s,30Phil.602;",
            "Mun. of Cavite v. Rojas",
            "",
            "",
            "30 Phil. 602",
        ),
        (
            "Cases: Yu v. De Lara, et.al. L16084, November 30, 1964",
            "Yu v. De Lara, et.al.",
            "L16084",
            "1964-11-30",
            "",
        ),
        (
            "Serg's Products, Inc. v. PCI Leasing and Finance Inc., 338 SCRA 499",
            "Serg's Products, Inc. v. PCI Leasing and Finance Inc.",
            "",
            "",
            "338 SCRA 499",
        ),
    ],
)
def test_irregular_syllabus_formats_remain_searchable(
    entry: str, title: str, docket: str, date: str, reporter: str
) -> None:
    case = parse_text(entry)[0]
    assert case.title == title
    assert case.docket == docket
    assert case.decision_date == date
    assert case.reporter_citation == reporter


def test_parse_numbered_week_list_preserves_assigned_title() -> None:
    cases = parse_text(SAMPLE)
    assert len(cases) == 3
    assert cases[0].week == "Week 2"
    assert cases[0].title == "Honda Philippines, Inc. v. Samahan ng Malalaya sa Honda"
    assert cases[0].docket == "G.R. No. 145561"
    assert cases[0].decision_date == "2005-06-15"
    assert cases[1].docket == "G.R. Nos. 158786 & 158789"
    assert cases[2].week == "Week 3"
    assert " v. " in cases[2].title
    assert not any(case.parse_error for case in cases)


def test_parse_wrapped_case_line() -> None:
    text = """WEEK 1
1. A Very Long Company Name v. Another Company,
G.R. No. 123456, January 2, 2020
"""
    cases = parse_text(text)
    assert len(cases) == 1
    assert cases[0].docket == "G.R. No. 123456"
    assert cases[0].decision_date == "2020-01-02"


def test_non_numbered_docx_paragraphs_remain_separate() -> None:
    document = Document()
    document.add_paragraph("WEEK 4")
    document.add_paragraph("First Case v. Second Case, G.R. No. 111111, March 4, 2019")
    document.add_paragraph("Third Case vs Fourth Case, G.R. No. 222222, April 5, 2020")
    stream = io.BytesIO()
    document.save(stream)
    text = extract_docx_text(stream.getvalue())
    cases = parse_text(text)
    assert len(cases) == 2
    assert [case.title for case in cases] == [
        "First Case v. Second Case",
        "Third Case v. Fourth Case",
    ]


def test_malformed_row_is_not_silently_dropped() -> None:
    cases = parse_text("WEEK 1\n1. Case without required information")
    assert len(cases) == 1
    assert cases[0].parse_error == ""
    assert "Missing or malformed docket number" in cases[0].warnings
    assert "Missing or malformed decision date" in cases[0].warnings


def test_l_number_is_supported() -> None:
    cases = parse_text(
        "1. PLDT v. NLRC, G.R. No. L-80609, August 23, 1988"
    )
    assert len(cases) == 1
    assert cases[0].docket == "G.R. No. L-80609"
    assert cases[0].decision_date == "1988-08-23"


def test_common_syllabus_citation_formats_are_separate_cases() -> None:
    cases = parse_text(
        """Mangila vs Court of Appeals (GR No 124027, 12 August 2002)
Mendiola vs Court of Appeals (497 SCRA 346)
JM Tuason & Co., Inc. vs Bolanos (95 Phil 106)"""
    )
    assert len(cases) == 3
    assert cases[0].title == "Mangila v. Court of Appeals"
    assert cases[0].docket == "GR No 124027"
    assert cases[0].decision_date == "2002-08-12"
    assert cases[1].reporter_citation == "497 SCRA 346"
    assert cases[2].reporter_citation == "95 Phil 106"
    assert not cases[1].parse_error


def test_syllabus_prose_is_ignored_and_module_is_preserved() -> None:
    cases = parse_text(
        """COURSE DESCRIPTION
This course surveys obligations and contracts.

MODULE 2
Assigned Readings
Mangila versus Court of Appeals (G.R. No. 124027, August 12, 2002)
Republic Act No. 386"""
    )
    assert len(cases) == 1
    assert cases[0].week == "Module 2"
    assert cases[0].title == "Mangila v. Court of Appeals"


def test_pdf_page_range_supports_single_pages_and_ranges() -> None:
    assert parse_page_range("2, 4-6", 8) == [1, 3, 4, 5]
    with pytest.raises(ValueError, match="between 1 and 8"):
        parse_page_range("9", 8)


def test_abbreviated_month_after_docket_without_comma_is_parsed() -> None:
    cases = parse_text("DENR Sec. v. Yap, GR No. 167707 Oct. 8, 2008")

    assert len(cases) == 1
    assert cases[0].title == "DENR Sec. v. Yap"
    assert cases[0].docket == "GR No. 167707"
    assert cases[0].decision_date == "2008-10-08"
    assert "Missing or malformed decision date" not in cases[0].warnings


@pytest.mark.parametrize(
    ("entry", "expected_title", "expected_clue"),
    [
        ("Lopez vs. Oroso Feb 28, 1958", "Lopez v. Oroso", "1958-02-28"),
        ("Tumalad v. Vicencio 122 SCRA 296", "Tumalad v. Vicencio", "122 SCRA 296"),
        (
            "Manila International Airport Authority v. CA (2006)",
            "Manila International Airport Authority v. CA",
            "2006",
        ),
        (
            "Province of Zamboanga Del Norte v. City of Zamboanga G.R. NO. L-24440",
            "Province of Zamboanga Del Norte v. City of Zamboanga",
            "L-24440",
        ),
        (
            "Case: People v. Pletcha, GR.R. No. 19029 June 27, 1977",
            "People v. Pletcha",
            "19029",
        ),
        (
            "Cases: Ayala Land v. Ray Burton 29 SCRA 48",
            "Ayala Land v. Ray Burton",
            "29 SCRA 48",
        ),
    ],
)
def test_previously_missed_benchmark_entries_keep_searchable_clues(
    entry: str, expected_title: str, expected_clue: str
) -> None:
    case = parse_text(entry)[0]
    assert case.title == expected_title
    assert expected_clue in " ".join(
        [case.docket, case.decision_date, case.reporter_citation, case.raw_text]
    )
