from __future__ import annotations

import uuid
from pathlib import Path

from openpyxl import load_workbook

from app.database import Database
from app.reporting import create_xlsx_report


def sample_case() -> dict[str, object]:
    return {
        "id": str(uuid.uuid4()),
        "ordinal": 1,
        "week": "Week 2",
        "title": "Honda Philippines, Inc. v. Samahan ng Malalaya sa Honda",
        "docket": "G.R. No. 145561",
        "decision_date": "2005-06-15",
        "raw_text": "sample",
        "filename": "Honda Philippines, Inc. v. Samahan ng Malalaya sa Honda.pdf",
        "warnings": [],
    }


def test_run_history_and_report_contain_not_saved_sheet(tmp_path: Path) -> None:
    db = Database(tmp_path / "history.sqlite3")
    case = sample_case()
    run_id = db.create_run("Test", str(tmp_path), [case])
    db.update_case(
        str(case["id"]),
        status="Not saved",
        error_code="NOT_FOUND",
        error_message="No matching case",
    )
    db.add_attempt(
        run_id,
        str(case["id"]),
        source="Lawphil",
        url="https://lawphil.net/test",
        attempt_no=1,
        status="Failed",
        error_code="NOT_FOUND",
        error_message="No matching case",
    )
    run = db.get_run(run_id)
    assert run is not None
    report = create_xlsx_report(run, tmp_path / "report.xlsx")
    workbook = load_workbook(report)
    assert workbook.sheetnames == ["Summary", "All Cases", "Not Saved", "Attempts"]
    not_saved = workbook["Not Saved"]
    assert not_saved.max_row == 2
    assert not_saved.cell(2, 11).value == "NOT_FOUND"


def test_interrupted_rows_are_recovered_on_database_open(tmp_path: Path) -> None:
    path = tmp_path / "history.sqlite3"
    db = Database(path)
    case = sample_case()
    run_id = db.create_run("Test", str(tmp_path), [case])
    db.update_run(run_id, status="Running")
    db.update_case(str(case["id"]), status="Searching Lawphil")
    recovered = Database(path).get_run(run_id)
    assert recovered is not None
    assert recovered["status"] == "Interrupted"
    assert recovered["cases"][0]["status"] == "Not saved"
    assert recovered["cases"][0]["error_code"] == "INTERRUPTED"
